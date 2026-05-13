import openpyxl
import psycopg2

# DB connection
conn = psycopg2.connect(
    host='aws-1-ap-southeast-1.pooler.supabase.com',
    port=6543,
    user='postgres.erjlvvznszwlxqdiduhq',
    password='ttH1TANuubY3JoMb',
    dbname='postgres',
    sslmode='require'
)
cur = conn.cursor()

# Get Vyttila warehouse ID
cur.execute("SELECT w.id, w.name, b.name FROM warehouses w JOIN branches b ON b.id = w.branch_id WHERE b.name ILIKE '%vyttila%'")
warehouses = cur.fetchall()
print('Vyttila warehouses:', warehouses)

if not warehouses:
    print('ERROR: No Vyttila warehouse found!')
    conn.close()
    exit()

wh_id = warehouses[0][0]
print('Using warehouse:', warehouses[0])
print()

# Get all stock for this warehouse from DB
cur.execute("""
    SELECT 
        v.variant_code,
        v.name as variant_name,
        v.price,
        v.cost_price,
        p.name as product_name,
        c.name as category_name,
        s.quantity
    FROM stocks s
    JOIN variants v ON v.id = s.variant_id
    JOIN products p ON p.id = v.product_id
    JOIN categories c ON c.id = p.category_id
    WHERE s.warehouse_id = %s
    ORDER BY c.name, v.variant_code
""", (wh_id,))
db_rows = cur.fetchall()
print('Total DB stock rows for Vyttila:', len(db_rows))

# Build DB lookup: variant_code -> dict
db_map = {}
for row in db_rows:
    code = row[0]
    db_map[code] = {
        'variant_name': row[1],
        'price': float(row[2]) if row[2] else 0,
        'cost_price': float(row[3]) if row[3] else 0,
        'product': row[4],
        'category': row[5],
        'qty': float(row[6]) if row[6] else 0
    }

# Parse Excel
excel_path = r'd:\QMark\defab_erp_backend\internal\migration\Defab Vyttila\LATEST_STOCK_MAY2026_WITH_GST.xlsx'
wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

excel_map = {}
excel_total_qty = 0
for sname in wb.sheetnames:
    ws = wb[sname]
    all_rows = list(ws.iter_rows(values_only=True))
    header_row = -1
    for i, row in enumerate(all_rows[:3]):
        if len(row) > 1 and row[1] is not None and str(row[1]).strip().upper() == 'CODE':
            header_row = i
            break
    if header_row < 0:
        continue
    sheet_count = 0
    for i, row in enumerate(all_rows):
        if i <= header_row or len(row) < 8:
            continue
        code = row[1]
        if code is None:
            continue
        try:
            code_int = int(float(str(code)))
        except:
            continue
        if code_int == 0:
            continue
        sp_excl = float(row[3]) if row[3] else 0
        cp = float(row[6]) if row[6] else 0
        qty = float(row[7]) if row[7] else 0
        item_desc = str(row[11]).strip().upper() if row[11] else sname.upper()
        if cp <= 0:
            cp = sp_excl
        if code_int in excel_map:
            excel_map[code_int]['qty'] += qty
        else:
            excel_map[code_int] = {
                'price': round(sp_excl, 2),
                'cp': round(cp, 2),
                'qty': qty,
                'category': sname,
                'item': item_desc
            }
        excel_total_qty += qty
        sheet_count += 1
    print(f'  Sheet: {sname} -> {sheet_count} rows')

wb.close()

print(f'\nExcel unique codes: {len(excel_map)}')
print(f'Excel total qty: {excel_total_qty}')
print()

# Compare
price_mismatches = []
cp_mismatches = []
qty_mismatches = []
missing_in_db = []
in_db_not_excel = []

for code, ex in excel_map.items():
    if code not in db_map:
        missing_in_db.append((code, ex['item'], ex['category']))
    else:
        db = db_map[code]
        # Price check (allow 0.1 rounding tolerance)
        if abs(ex['price'] - db['price']) > 0.1:
            price_mismatches.append((code, ex['price'], db['price'], ex['item']))
        if abs(ex['cp'] - db['cost_price']) > 0.1:
            cp_mismatches.append((code, ex['cp'], db['cost_price'], ex['item']))
        if abs(ex['qty'] - db['qty']) > 0.01:
            qty_mismatches.append((code, ex['qty'], db['qty'], ex['item']))

for code in db_map:
    if code not in excel_map:
        db = db_map[code]
        in_db_not_excel.append((code, db['product'], db['category']))

print('=' * 60)
print(f'Missing in DB (in Excel but not in DB): {len(missing_in_db)}')
print(f'Extra in DB (in DB but not in Excel): {len(in_db_not_excel)}')
print(f'Price mismatches: {len(price_mismatches)}')
print(f'CP mismatches: {len(cp_mismatches)}')
print(f'QTY mismatches: {len(qty_mismatches)}')
print('=' * 60)

if missing_in_db:
    print('\nMissing in DB:')
    for c, item, cat in missing_in_db[:20]:
        print(f'  code={c}  item={item}  sheet={cat}')

if in_db_not_excel:
    print('\nExtra in DB (not in Excel):')
    for c, prod, cat in in_db_not_excel[:20]:
        print(f'  code={c}  product={prod}  category={cat}')

if price_mismatches:
    print('\nPrice mismatches (Excel vs DB):')
    for c, ep, dp, item in price_mismatches[:20]:
        print(f'  code={c}  item={item}  excel={ep}  db={dp}  diff={round(ep-dp,2)}')

if cp_mismatches:
    print('\nCP mismatches:')
    for c, ep, dp, item in cp_mismatches[:20]:
        print(f'  code={c}  item={item}  excel_cp={ep}  db_cp={dp}  diff={round(ep-dp,2)}')

if qty_mismatches:
    print('\nQTY mismatches:')
    for c, eq, dq, item in qty_mismatches[:20]:
        print(f'  code={c}  item={item}  excel_qty={eq}  db_qty={dq}  diff={round(eq-dq,2)}')

# Summary
all_good = (len(missing_in_db) == 0 and len(price_mismatches) == 0 and
            len(cp_mismatches) == 0 and len(qty_mismatches) == 0)
if all_good:
    print('\n✅ ALL DATA MATCHES PERFECTLY!')
else:
    print(f'\n⚠️  Issues found — review above')

cur.close()
conn.close()
